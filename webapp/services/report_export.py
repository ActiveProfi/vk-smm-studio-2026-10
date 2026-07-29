"""Экспорт месячного отчёта в Excel и PDF."""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_FONT_REGISTERED = False
_FONT_NAME = "Helvetica"


def _register_cyrillic_font() -> str:
    """Подключаем системный шрифт с кириллицей (Windows Arial)."""
    global _FONT_REGISTERED, _FONT_NAME
    if _FONT_REGISTERED:
        return _FONT_NAME
    candidates = [
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\arialuni.ttf"),
        Path(r"C:\Windows\Fonts\segoeui.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for path in candidates:
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont("ReportCyr", str(path)))
                _FONT_NAME = "ReportCyr"
                _FONT_REGISTERED = True
                return _FONT_NAME
            except Exception:  # noqa: BLE001
                continue
    _FONT_REGISTERED = True
    return _FONT_NAME


def safe_filename(name: str) -> str:
    cleaned = re.sub(r"[^\w\-а-яА-ЯёЁ]+", "_", (name or "channel").strip(), flags=re.UNICODE)
    return (cleaned[:60] or "channel").strip("_")


def build_excel_bytes(report: dict[str, Any], *, channel_name: str, month_title: str) -> bytes:
    wb = Workbook()

    # --- Сводка ---
    ws1 = wb.active
    ws1.title = "Сводка"
    ws1["A1"] = "Месячный отчёт клиенту"
    ws1["A1"].font = Font(bold=True, size=14)
    ws1["A2"] = f"Канал: {channel_name}"
    ws1["A3"] = f"Период: {month_title} ({report.get('period_label') or ''})"
    pkg = report.get("package") or {}
    ws1["A4"] = (
        f"Пакет: {pkg.get('title') or '—'} · использовано {pkg.get('used', 0)}"
        + (f"/{pkg.get('limit')}" if not pkg.get("unlimited") else " (без лимита)")
    )

    counts = report.get("counts") or {}
    ws1["A6"] = "Показатель"
    ws1["B6"] = "Значение"
    ws1["A6"].font = Font(bold=True)
    ws1["B6"].font = Font(bold=True)
    rows = [
        ("Всего постов", counts.get("total", 0)),
        ("В плане", counts.get("planned", 0)),
        ("Готово", counts.get("ready", 0)),
        ("Опубликовано", counts.get("published", 0)),
        ("На согласовании", counts.get("pending_approval", 0)),
        ("Отклонено", counts.get("rejected", 0)),
        ("Расходы API, ₽", (report.get("costs") or {}).get("total_rub", 0)),
    ]
    for i, (label, value) in enumerate(rows, start=7):
        ws1[f"A{i}"] = label
        ws1[f"B{i}"] = value

    ws1["A15"] = "По рубрикам"
    ws1["A15"].font = Font(bold=True)
    ws1["A16"] = "Рубрика"
    ws1["B16"] = "Кол-во"
    ws1["A16"].font = Font(bold=True)
    ws1["B16"].font = Font(bold=True)
    r = 17
    for title, cnt in report.get("by_rubric") or []:
        ws1[f"A{r}"] = title
        ws1[f"B{r}"] = cnt
        r += 1

    ws1.column_dimensions["A"].width = 36
    ws1.column_dimensions["B"].width = 18

    # --- Посты ---
    ws2 = wb.create_sheet("Посты")
    headers = ["Когда", "Тема", "Рубрика", "Статус", "Клиент", "VK post id"]
    for col, h in enumerate(headers, start=1):
        cell = ws2.cell(1, col, h)
        cell.font = Font(bold=True)
    for row_idx, post in enumerate(report.get("posts") or [], start=2):
        ws2.cell(row_idx, 1, post.get("when") or "")
        ws2.cell(row_idx, 2, post.get("topic") or "")
        ws2.cell(row_idx, 3, post.get("rubric") or "")
        ws2.cell(row_idx, 4, post.get("status") or "")
        ws2.cell(row_idx, 5, post.get("client_status_label") or "")
        ws2.cell(row_idx, 6, post.get("vk_post_id") or "")
        ws2.cell(row_idx, 2).alignment = Alignment(wrap_text=True)
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 48
    ws2.column_dimensions["C"].width = 22
    ws2.column_dimensions["D"].width = 20
    ws2.column_dimensions["E"].width = 16
    ws2.column_dimensions["F"].width = 14

    # --- Расходы ---
    ws3 = wb.create_sheet("Расходы API")
    ws3["A1"] = "Вид"
    ws3["B1"] = "Кол-во"
    ws3["C1"] = "Сумма, ₽"
    for col in range(1, 4):
        ws3.cell(1, col).font = Font(bold=True)
    for row_idx, row in enumerate(report.get("cost_rows") or [], start=2):
        ws3.cell(row_idx, 1, row.get("title") or row.get("kind") or "")
        ws3.cell(row_idx, 2, row.get("count") or 0)
        ws3.cell(row_idx, 3, round(float(row.get("cost") or 0), 2))
    last = 2 + len(report.get("cost_rows") or [])
    ws3.cell(last, 1, "Итого")
    ws3.cell(last, 1).font = Font(bold=True)
    ws3.cell(last, 3, (report.get("costs") or {}).get("total_rub", 0))
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 12
    ws3.column_dimensions["C"].width = 14

    # Журнал операций
    ws4 = wb.create_sheet("Журнал API")
    for col, h in enumerate(["Когда", "Тип", "Сумма, ₽", "Комментарий"], start=1):
        cell = ws4.cell(1, col, h)
        cell.font = Font(bold=True)
    for row_idx, ev in enumerate(report.get("cost_events") or [], start=2):
        ws4.cell(row_idx, 1, ev.get("when") or "")
        ws4.cell(row_idx, 2, ev.get("title") or ev.get("kind") or "")
        ws4.cell(row_idx, 3, ev.get("cost") or 0)
        ws4.cell(row_idx, 4, ev.get("note") or "")
    ws4.column_dimensions["A"].width = 18
    ws4.column_dimensions["B"].width = 24
    ws4.column_dimensions["C"].width = 12
    ws4.column_dimensions["D"].width = 40

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf_bytes(report: dict[str, Any], *, channel_name: str, month_title: str) -> bytes:
    font = _register_cyrillic_font()
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"Отчёт · {channel_name} · {month_title}",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontName=font,
        fontSize=16,
        leading=20,
        spaceAfter=6,
    )
    body = ParagraphStyle(
        "ReportBody",
        parent=styles["Normal"],
        fontName=font,
        fontSize=10,
        leading=14,
        spaceAfter=4,
    )
    h2 = ParagraphStyle(
        "ReportH2",
        parent=styles["Heading2"],
        fontName=font,
        fontSize=12,
        leading=16,
        spaceBefore=10,
        spaceAfter=6,
    )

    story: list[Any] = []
    story.append(Paragraph("Месячный отчёт клиенту", title_style))
    story.append(Paragraph(f"Канал: <b>{_esc(channel_name)}</b>", body))
    story.append(Paragraph(f"Период: {_esc(month_title)} ({_esc(report.get('period_label') or '')})", body))
    pkg = report.get("package") or {}
    pkg_line = f"Пакет: {_esc(pkg.get('title') or '—')} · использовано {pkg.get('used', 0)}"
    if not pkg.get("unlimited"):
        pkg_line += f"/{pkg.get('limit')}"
    else:
        pkg_line += " (без лимита)"
    story.append(Paragraph(pkg_line, body))
    story.append(Spacer(1, 6))

    counts = report.get("counts") or {}
    summary_data = [
        ["Показатель", "Значение"],
        ["Всего постов", str(counts.get("total", 0))],
        ["В плане", str(counts.get("planned", 0))],
        ["Готово", str(counts.get("ready", 0))],
        ["Опубликовано", str(counts.get("published", 0))],
        ["На согласовании", str(counts.get("pending_approval", 0))],
        ["Отклонено", str(counts.get("rejected", 0))],
        ["Расходы API, ₽", str((report.get("costs") or {}).get("total_rub", 0))],
    ]
    story.append(Paragraph("Сводка", h2))
    story.append(_table(summary_data, font, col_widths=[90 * mm, 60 * mm]))

    rubric_rows = [["Рубрика", "Кол-во"]]
    for title, cnt in report.get("by_rubric") or []:
        rubric_rows.append([_esc(str(title)), str(cnt)])
    if len(rubric_rows) == 1:
        rubric_rows.append(["Нет данных", "0"])
    story.append(Paragraph("По рубрикам", h2))
    story.append(_table(rubric_rows, font, col_widths=[120 * mm, 30 * mm]))

    post_rows = [["Когда", "Тема", "Статус", "Клиент"]]
    for post in report.get("posts") or []:
        post_rows.append(
            [
                _esc(post.get("when") or "—"),
                _esc(post.get("topic") or "—"),
                _esc(post.get("status") or "—"),
                _esc(post.get("client_status_label") or "—"),
            ]
        )
    if len(post_rows) == 1:
        post_rows.append(["—", "Нет постов за месяц", "—", "—"])
    story.append(Paragraph("Посты месяца", h2))
    story.append(_table(post_rows, font, col_widths=[40 * mm, 70 * mm, 35 * mm, 25 * mm], font_size=8))

    cost_rows = [["Вид", "Кол-во", "Сумма, ₽"]]
    for row in report.get("cost_rows") or []:
        cost_rows.append(
            [
                _esc(str(row.get("title") or row.get("kind") or "")),
                str(row.get("count") or 0),
                f"{float(row.get('cost') or 0):.2f}",
            ]
        )
    cost_rows.append(["Итого", "", str((report.get("costs") or {}).get("total_rub", 0))])
    story.append(Paragraph("Расходы API (оценка)", h2))
    story.append(_table(cost_rows, font, col_widths=[60 * mm, 40 * mm, 50 * mm]))

    doc.build(story)
    return buf.getvalue()


def _esc(text: str) -> str:
    return (
        (text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _table(data: list[list[str]], font: str, *, col_widths: list[float], font_size: int = 9) -> Table:
    styled: list[list[Any]] = []
    for row in data:
        styled.append(
            [
                Paragraph(str(cell), ParagraphStyle("td", fontName=font, fontSize=font_size, leading=font_size + 3))
                for cell in row
            ]
        )
    table = Table(styled, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("FONTSIZE", (0, 0), (-1, -1), font_size),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table
