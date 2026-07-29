"""Извлечение текста из загруженных материалов: txt, md, Word, PDF, Excel."""
from __future__ import annotations

from pathlib import Path

ALLOWED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".docx",
    ".doc",
    ".pdf",
    ".xlsx",
    ".xls",
}


def is_allowed_filename(name: str) -> bool:
    lower = (name or "").lower()
    return any(lower.endswith(ext) for ext in ALLOWED_EXTENSIONS)


def file_extension(name: str) -> str:
    lower = (name or "").lower()
    for ext in sorted(ALLOWED_EXTENSIONS, key=len, reverse=True):
        if lower.endswith(ext):
            return ext
    return Path(name or "").suffix.lower()


def safe_storage_name(original_name: str) -> str:
    """
    Безопасное имя для диска с сохранением расширения.
    werkzeug.secure_filename съедает кириллицу («Материалы.docx» → «docx»),
    из‑за этого Word-файлы ошибочно отклонялись.
    """
    from werkzeug.utils import secure_filename

    original = (original_name or "").strip() or "material.txt"
    ext = file_extension(original)
    stem = Path(original).stem
    safe_stem = secure_filename(stem) or ""
    # если имя было только кириллицей — stem пустой или совпал с «docx»
    if not safe_stem or (ext and safe_stem.lower() == ext.lstrip(".").lower()):
        safe_stem = "material"
    if not ext:
        ext = ".txt"
    return f"{safe_stem}{ext}"


def companion_path(path: Path) -> Path:
    return Path(str(path) + ".extracted.txt")


def extract_text(path: Path | str) -> str:
    path = Path(path)
    if not path.exists():
        return ""
    suffix = path.suffix.lower()

    try:
        if suffix in {".txt", ".md", ".markdown", ".csv"}:
            return _read_plain(path)
        if suffix == ".docx":
            return _read_docx(path)
        if suffix == ".doc":
            return _read_doc_legacy(path)
        if suffix == ".pdf":
            return _read_pdf(path)
        if suffix == ".xlsx":
            return _read_xlsx(path)
        if suffix == ".xls":
            return _read_xls(path)
    except Exception as exc:  # noqa: BLE001
        return f"[Не удалось прочитать файл {path.name}: {exc}]"

    return f"[Формат {suffix or 'без расширения'} пока не поддерживается]"


def extract_and_cache(path: Path | str) -> str:
    """Извлекает текст и сохраняет рядом служебный .extracted.txt для саммари."""
    path = Path(path)
    text = extract_text(path).strip()
    companion_path(path).write_text(text, encoding="utf-8")
    return text


def read_cached_or_extract(path: Path | str) -> str:
    path = Path(path)
    cached = companion_path(path)
    if cached.exists():
        return cached.read_text(encoding="utf-8", errors="replace").strip()
    return extract_and_cache(path)


def _read_plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            line = " | ".join(x for x in cells if x)
            if line:
                parts.append(line)
    return "\n".join(parts)


def _read_doc_legacy(path: Path) -> str:
    """
    Старый .doc (Word 97–2003).
    Пробуем через Word на Windows, иначе подсказка сохранить как .docx.
    """
    try:
        import win32com.client  # type: ignore

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(path.resolve()))
        try:
            text = doc.Content.Text
        finally:
            doc.Close(False)
            word.Quit()
        return (text or "").replace("\r", "\n").strip()
    except Exception:
        return (
            f"[Файл {path.name} в старом формате .doc. "
            "Сохраните его как .docx и загрузите снова — так надёжнее.]"
        )


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        t = t.strip()
        if t:
            parts.append(f"--- страница {i} ---\n{t}")
    return "\n\n".join(parts) if parts else f"[В PDF {path.name} не найден текстовый слой]"


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"=== Лист: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _read_xls(path: Path) -> str:
    import xlrd

    book = xlrd.open_workbook(str(path))
    parts: list[str] = []
    for sheet in book.sheets():
        parts.append(f"=== Лист: {sheet.name} ===")
        for r in range(sheet.nrows):
            cells = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)
